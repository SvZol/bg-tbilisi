"""Штрафы в результатах и автозакрытие регистрации."""
import io
from datetime import timedelta

import openpyxl

from core.registration import now_tbilisi
from models.content import EventQuestion
from models.event import Event
from models.team import Team
from tests.conftest import auth_headers, make_user


def make_team(client, headers, event_id, name, category="adult"):
    r = client.post("/teams/", headers=headers, json={"event_id": event_id, "name": name, "category": category,
                                                      "member_count": 2 if category == "child" else 1})
    assert r.status_code == 200, r.text
    return r.json()["id"]


def make_questions(db, event_id):
    for number, kind in ((1, "КП"), (101, "Задача"), (2, "КП"), (102, "Задача")):
        db.add(EventQuestion(event_id=event_id, number=number, kp_type=kind, text="?", correct_answer="x"))
    db.commit()


def results_xlsx(team_names, points, penalties=None, file_sums=None):
    """Файл в формате «проверка»: points[team] = (задание КП-1, задача КП-1, задание КП-2, задача КП-2)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "проверка"
    ws.cell(5, 1, "КП-00")
    ws.cell(5, 3, "Эталон-команда")
    for i, name in enumerate(team_names):
        ws.cell(5, 4 + i, name)
    layout = [(7, "КП-01", "задание КП", 0), (8, None, "задача", 1), (9, "КП-02", "задание КП", 2), (10, None, "задача", 3)]
    for row, label, kind, idx in layout:
        ws.cell(row, 1, label)
        ws.cell(row, 2, kind)
        ws.cell(row, 3, "x")
        for i, name in enumerate(team_names):
            ws.cell(row, 4 + i, points[name][idx])
    ws.cell(12, 3, "КП")
    ws.cell(13, 3, "задачи")
    ws.cell(14, 3, "штраф")
    ws.cell(15, 3, "Сумма")
    for i, name in enumerate(team_names):
        p = points[name]
        ws.cell(12, 4 + i, p[0] + p[2])
        ws.cell(13, 4 + i, p[1] + p[3])
        pen = (penalties or {}).get(name)
        if pen is not None:
            ws.cell(14, 4 + i, pen)
        ws.cell(15, 4 + i, (file_sums or {}).get(name, p[0] + p[1] + p[2] + p[3] + (pen or 0)))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def upload(client, headers, event_id, content):
    return client.post(f"/admin/events/{event_id}/import-results-excel", headers=headers,
                       files={"file": ("r.xlsx", content, "application/octet-stream")})


def setup(client, db, admin_h, event_id):
    make_user(db, "t1@x.com")
    make_user(db, "t2@x.com")
    h1, h2 = auth_headers(client, "t1@x.com"), auth_headers(client, "t2@x.com")
    make_team(client, h1, event_id, "Альфа")
    make_team(client, h2, event_id, "Бета")
    make_questions(db, event_id)
    return ["Альфа", "Бета"]


POINTS = {"Альфа": (1, 1, 1, 1), "Бета": (1, 1, 1, 0)}


def scoreboard(client, admin_h, event_id):
    rows = client.get(f"/admin/events/{event_id}/scoreboard", headers=admin_h).json()["adult"]
    return {r["name"]: r for r in rows}


def test_penalty_is_subtracted_from_total(client, db, admin_h, event_id):
    names = setup(client, db, admin_h, event_id)
    r = upload(client, admin_h, event_id, results_xlsx(names, POINTS, penalties={"Альфа": -3}))
    assert r.status_code == 200 and r.json()["warnings"] == []
    board = scoreboard(client, admin_h, event_id)
    assert board["Альфа"]["score"] == 4 - 3 and board["Альфа"]["penalty"] == -3
    assert board["Бета"]["score"] == 3 and board["Бета"]["penalty"] == 0
    assert board["Бета"]["rank"] == 1 and board["Альфа"]["rank"] == 2   # штраф меняет места


def test_positive_penalty_still_subtracts(client, db, admin_h, event_id):
    names = setup(client, db, admin_h, event_id)
    upload(client, admin_h, event_id, results_xlsx(names, POINTS, penalties={"Альфа": 2}))
    assert scoreboard(client, admin_h, event_id)["Альфа"]["score"] == 4 - 2


def test_reimport_without_penalty_clears_old_penalty(client, db, admin_h, event_id):
    names = setup(client, db, admin_h, event_id)
    upload(client, admin_h, event_id, results_xlsx(names, POINTS, penalties={"Альфа": -3}))
    upload(client, admin_h, event_id, results_xlsx(names, POINTS))
    assert scoreboard(client, admin_h, event_id)["Альфа"]["score"] == 4


def test_warns_when_file_sum_differs_from_ours(client, db, admin_h, event_id):
    names = setup(client, db, admin_h, event_id)
    r = upload(client, admin_h, event_id, results_xlsx(names, POINTS, file_sums={"Бета": 99}))
    assert len(r.json()["warnings"]) == 1 and "Бета" in r.json()["warnings"][0]


def test_summary_rows_are_not_imported_as_results(client, db, admin_h, event_id):
    names = setup(client, db, admin_h, event_id)
    assert upload(client, admin_h, event_id, results_xlsx(names, POINTS)).json()["created"] == 8


def test_public_page_data_includes_penalty(client, db, admin_h, event_id):
    names = setup(client, db, admin_h, event_id)
    upload(client, admin_h, event_id, results_xlsx(names, POINTS, penalties={"Бета": -1}))
    client.post(f"/admin/events/{event_id}/publish-results", headers=admin_h)
    teams = {t["name"]: t for t in client.get(f"/admin/public/events/{event_id}/questions").json()["teams"]}
    assert teams["Бета"]["penalty"] == -1 and teams["Альфа"]["penalty"] == 0
    team_id = db.query(Team).filter(Team.name == "Бета").first().id
    assert client.get(f"/teams/{team_id}/public").json()["penalty"] == -1


def test_results_email_contains_total_with_penalty(client, db, admin_h, event_id, outbox):
    names = setup(client, db, admin_h, event_id)
    upload(client, admin_h, event_id, results_xlsx(names, POINTS, penalties={"Альфа": -3}))
    outbox.clear()
    client.post(f"/admin/events/{event_id}/publish-results", headers=admin_h)
    alpha = next(m for m in outbox if m[0] == "t1@x.com")
    assert "баллы: 1" in alpha[2]


def test_results_template_has_summary_rows(client, db, admin_h, event_id):
    setup(client, db, admin_h, event_id)
    r = client.get(f"/admin/events/{event_id}/results-template", headers=admin_h)
    ws = openpyxl.load_workbook(io.BytesIO(r.content)).active
    labels = [c.value for c in ws["C"] if c.value in ("КП", "задачи", "штраф", "Сумма")]
    assert labels == ["КП", "задачи", "штраф", "Сумма"]


# --- автозакрытие регистрации ---

def set_deadline(db, event_id, delta):
    event = db.query(Event).filter(Event.id == event_id).first()
    event.reg_deadline = now_tbilisi() + delta
    db.commit()


def test_registration_open_before_deadline(client, db, user_h, event_id):
    set_deadline(db, event_id, timedelta(hours=1))
    assert client.get(f"/events/{event_id}").json()["status"] == "open"
    assert client.post("/teams/", headers=user_h, json={"event_id": event_id, "name": "Ранние"}).status_code == 200


def test_registration_closes_itself_after_deadline(client, db, user_h, event_id):
    set_deadline(db, event_id, -timedelta(minutes=1))
    assert client.get(f"/events/{event_id}").json()["status"] == "closed"
    assert client.get("/events/").json()[0]["status"] == "closed"
    r = client.post("/teams/", headers=user_h, json={"event_id": event_id, "name": "Поздние"})
    assert r.status_code == 400


def test_admin_still_sees_real_status(client, db, admin_h, event_id):
    set_deadline(db, event_id, -timedelta(hours=1))
    assert client.get("/admin/events", headers=admin_h).json()[0]["status"] == "open"


def test_finished_event_stays_finished(client, db, admin_h, event_id):
    client.patch(f"/admin/events/{event_id}/status?status=finished", headers=admin_h)
    set_deadline(db, event_id, -timedelta(days=5))
    assert client.get(f"/events/{event_id}").json()["status"] == "finished"


def test_publish_sends_emails_only_once(client, db, admin_h, event_id, outbox):
    names = setup(client, db, admin_h, event_id)
    upload(client, admin_h, event_id, results_xlsx(names, POINTS))
    outbox.clear()
    first = client.post(f"/admin/events/{event_id}/publish-results", headers=admin_h).json()
    assert first["emails_sent"] and len(outbox) == 2
    outbox.clear()
    second = client.post(f"/admin/events/{event_id}/publish-results", headers=admin_h).json()
    assert not second["emails_sent"] and second["published"] == 4 and outbox == []


def test_publish_after_hide_sends_emails_again(client, db, admin_h, event_id, outbox):
    names = setup(client, db, admin_h, event_id)
    upload(client, admin_h, event_id, results_xlsx(names, POINTS))
    client.post(f"/admin/events/{event_id}/publish-results", headers=admin_h)
    client.post(f"/admin/events/{event_id}/unpublish-results", headers=admin_h)
    outbox.clear()
    assert client.post(f"/admin/events/{event_id}/publish-results", headers=admin_h).json()["emails_sent"]
