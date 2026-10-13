"""Места при равных баллах, время финиша и адреса КП."""
import io
from datetime import datetime, time

import openpyxl

from core.scoring import parse_finish_time, rank_teams
from models.content import EventQuestion
from models.team import Team
from tests.conftest import auth_headers, make_user
from tests.test_scoring import make_team, make_questions, upload


class FakeTeam:
    def __init__(self, tid):
        self.id = tid


def ranks(rows):
    """rows: {id: (total, finish_time)} -> {id: place}"""
    teams = [FakeTeam(i) for i in rows]
    scores = {str(i): {"total": t, "finish_time": f} for i, (t, f) in rows.items()}
    return {int(k): v for k, v in rank_teams(teams, scores).items()}


def test_equal_points_and_time_share_place():
    r = ranks({1: (41, None), 2: (38, "01:00:00"), 3: (38, "01:00:00"), 4: (38, "01:00:00"), 5: (37, None)})
    assert [r[i] for i in (1, 2, 3, 4, 5)] == [1, 2, 2, 2, 5]


def test_earlier_finish_wins_a_tie():
    r = ranks({1: (30, "01:10:00"), 2: (30, "01:05:00"), 3: (29, "00:50:00")})
    assert (r[2], r[1], r[3]) == (1, 2, 3)


def test_equal_points_without_time_are_a_tie():
    r = ranks({1: (30, None), 2: (30, None)})
    assert r[1] == r[2] == 1


def test_team_with_time_beats_team_without_time_on_equal_points():
    r = ranks({1: (30, None), 2: (30, "01:00:00")})
    assert (r[2], r[1]) == (1, 2)


def test_more_points_beats_earlier_time():
    r = ranks({1: (31, "02:00:00"), 2: (30, "00:30:00")})
    assert (r[1], r[2]) == (1, 2)


def test_parse_finish_time_formats():
    assert parse_finish_time(time(1, 5, 9)) == "01:05:09"
    assert parse_finish_time(datetime(2026, 10, 11, 15, 20)) == "15:20:00"
    assert parse_finish_time(0.5) == "12:00:00"
    assert parse_finish_time("1:05") == "01:05:00"
    assert parse_finish_time(" 01:05:09 ") == "01:05:09"
    for bad in (None, "", "скоро", "25:99", 5, True):
        assert parse_finish_time(bad) is None


# --- импорт результатов с временем финиша ---

def results_file(names, points, times):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "проверка"
    ws.cell(5, 3, "Эталон-команда")
    for i, n in enumerate(names):
        ws.cell(5, 4 + i, n)
    for row, label, kind, idx in [(7, "КП-01", "задание КП", 0), (8, None, "задача", 1), (9, "КП-02", "задание КП", 2), (10, None, "задача", 3)]:
        ws.cell(row, 1, label); ws.cell(row, 2, kind)
        for i, n in enumerate(names):
            ws.cell(row, 4 + i, points[n][idx])
    ws.cell(66, 2, "Время финиша")
    for i, n in enumerate(names):
        if times.get(n) is not None:
            ws.cell(66, 4 + i, times[n])
    buf = io.BytesIO(); wb.save(buf)
    return buf.getvalue()


def three_teams(client, db, event_id):
    names = []
    for i in range(3):
        make_user(db, f"t{i}@x.com")
        name = f"Команда{i}"
        make_team(client, auth_headers(client, f"t{i}@x.com"), event_id, name)
        names.append(name)
    make_questions(db, event_id)
    return names


def board(client, admin_h, event_id):
    return {r["name"]: r for r in client.get(f"/admin/events/{event_id}/scoreboard", headers=admin_h).json()["adult"]}


def test_import_ties_share_place_and_kp_points_do_not_split_them(client, db, admin_h, event_id):
    a, b, c = three_teams(client, db, event_id)
    # одинаковый итог 4, но очки за КП/задачи распределены по-разному
    pts = {a: (1, 1, 1, 1), b: (1, 1, 1, 1), c: (1, 1, 0, 2)}
    assert upload(client, admin_h, event_id, results_file([a, b, c], pts, {})).status_code == 200
    got = board(client, admin_h, event_id)
    assert got[a]["rank"] == got[b]["rank"] == got[c]["rank"] == 1


def test_import_finish_time_breaks_tie(client, db, admin_h, event_id):
    a, b, c = three_teams(client, db, event_id)
    pts = {n: (1, 1, 1, 1) for n in (a, b, c)}
    times = {a: time(1, 30), b: time(1, 10), c: time(1, 30)}
    upload(client, admin_h, event_id, results_file([a, b, c], pts, times))
    got = board(client, admin_h, event_id)
    assert (got[b]["rank"], got[a]["rank"], got[c]["rank"]) == (1, 2, 2)
    assert got[b]["finish_time"] == "01:10:00"


def test_reimport_without_times_clears_them(client, db, admin_h, event_id):
    a, b, c = three_teams(client, db, event_id)
    pts = {n: (1, 1, 1, 1) for n in (a, b, c)}
    upload(client, admin_h, event_id, results_file([a, b, c], pts, {a: time(1, 0)}))
    upload(client, admin_h, event_id, results_file([a, b, c], pts, {}))
    assert board(client, admin_h, event_id)[a]["finish_time"] is None


def test_public_page_gets_ranks_from_server(client, db, admin_h, event_id):
    a, b, c = three_teams(client, db, event_id)
    pts = {a: (1, 1, 1, 1), b: (1, 1, 1, 1), c: (0, 0, 0, 0)}
    upload(client, admin_h, event_id, results_file([a, b, c], pts, {}))
    client.post(f"/admin/events/{event_id}/publish-results", headers=admin_h)
    teams = {t["name"]: t for t in client.get(f"/admin/public/events/{event_id}/questions").json()["teams"]}
    assert (teams[a]["rank"], teams[b]["rank"], teams[c]["rank"]) == (1, 1, 3)
    assert teams[a]["total"] == 4 and teams[a]["kp"] == 2 and teams[a]["tasks"] == 2


def test_template_has_finish_time_row(client, db, admin_h, event_id):
    three_teams(client, db, event_id)
    r = client.get(f"/admin/events/{event_id}/results-template", headers=admin_h)
    ws = openpyxl.load_workbook(io.BytesIO(r.content)).active
    assert "Время финиша" in [c.value for c in ws["B"]]


# --- адреса КП ---

def kp_file():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Все КП"
    ws.append(["дата", "тип КП", "номер п/п", "КП (старый)", "номер КП", "Адрес (рус)", "Адрес (англ)", "Объект", "Задание", "Ответ", "Комментарий", "Задача", "Ответ задачи"])
    ws.append([None, "Старт", None, None, "КП-00", "Парк Рике", None, "Амфитеатр", "Старт игры"])
    ws.append([None, "Адрес", None, None, "КП-01", "Ул. Гришашвили, 54", "Grishashvili St", "Памятник", "Сумма цифр", 36, None, "Задачка", 12])
    ws.append([None, "ФотоКП", None, None, "КП-42", "II пер. Горгасала", "Gorgasli Dead End", "Табличка", "Название улицы", "II пер. Горгасала"])
    buf = io.BytesIO(); wb.save(buf)
    return buf.getvalue()


def import_kp(client, headers, event_id):
    return client.post(f"/admin/events/{event_id}/import-kp-excel", headers=headers,
                       files={"file": ("kp.xlsx", kp_file(), "application/octet-stream")})


def test_kp_import_stores_addresses(client, db, admin_h, event_id):
    assert import_kp(client, admin_h, event_id).status_code == 200
    q = {x.number: x for x in db.query(EventQuestion).filter(EventQuestion.event_id == event_id)}
    assert (q[1].address_ru, q[1].address_en, q[1].object_note) == ("Ул. Гришашвили, 54", "Grishashvili St", "Памятник")
    assert q[42].address_ru == "II пер. Горгасала" and q[42].kp_type == "ФотоКП"
    assert q[0].address_ru == "Парк Рике" and q[0].address_en is None
    assert q[101].address_ru is None  # адрес хранится только у задания


def test_kp_reimport_fills_addresses_of_existing_questions(client, db, admin_h, event_id):
    db.add(EventQuestion(event_id=event_id, number=1, kp_type="Адрес", text="Сумма цифр", correct_answer="36"))
    db.commit()
    import_kp(client, admin_h, event_id)
    db.expire_all()
    assert db.query(EventQuestion).filter(EventQuestion.event_id == event_id, EventQuestion.number == 1).first().address_ru == "Ул. Гришашвили, 54"


def test_addresses_hidden_until_published(client, db, admin_h, event_id):
    import_kp(client, admin_h, event_id)
    assert client.get(f"/admin/public/events/{event_id}/questions").json()["questions"] == []
    client.post(f"/admin/events/{event_id}/publish-results", headers=admin_h)
    qs = client.get(f"/admin/public/events/{event_id}/questions").json()["questions"]
    assert "Ул. Гришашвили, 54" in [q["address_ru"] for q in qs]
